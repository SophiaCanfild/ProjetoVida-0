#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
VIDA+ — Gerador do TCC em .docx seguindo as normas ABNT
========================================================

Converte o arquivo de texto  docs/tcc/TCC-Vida+.md  no documento
docs/tcc/TCC-Vida+.docx  já formatado conforme:

  • ABNT NBR 14724:2024 (versão corrigida 01.04.2025) — apresentação de
    trabalhos acadêmicos: formato, margens, espaçamento, paginação,
    numeração progressiva, ilustrações e tabelas;
  • ABNT NBR 6024:2012 — numeração progressiva das seções;
  • ABNT NBR 6027:2012 — sumário;
  • ABNT NBR 6028:2021 — resumo / palavras-chave;
  • ABNT NBR 10520:2023 — citações;
  • ABNT NBR 6023 (edição vigente) — referências.

COMO USAR
---------
    python3 docs/tcc/gerar-docx.py

Requisitos: python-docx e Pillow  (pip install python-docx Pillow)

SINTAXE ACEITA NO ARQUIVO .md
-----------------------------
  # TÍTULO ............ seção primária (começa em nova página)
  ## Título .......... seção secundária
  ### Título ......... seção terciária (itálico)
  #### Título ........ seção quaternária (itálico + negrito)
  > texto ............ citação direta com mais de 3 linhas
  | a | b | .......... tabela (linha separadora |---| obrigatória)
  ![título](arquivo) . ilustração (centralizada, largura máx. 15 cm)
  [[ ... ]] .......... marcação de PENDÊNCIA (sai em vermelho itálico)
  <!-- ... --> ....... comentário: NÃO aparece no .docx (é só orientação)
  <!-- SUMARIO --> ... insere o campo automático de SUMÁRIO do Word

  Parágrafos com posicionamento especial usam etiquetas no início:
    [e]              linha em branco
    [c]   [cb]       centralizado / centralizado em negrito
    [r]   [rb]       à direita / à direita em negrito
    [n]              bloco "natureza do trabalho", dedicatória e epígrafe
    [s]              sem recuo de primeira linha (resumo e palavras-chave)
    [l]              item de lista com pontilhado — "texto | página"
    [a]              alínea a) b) c) com recuo próprio (NBR 6024:2012)
    [c14] [cb14]     idem, com fonte 14 (títulos de capa)
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ------------------------------------------------------------------
# Constantes de formatação (ABNT NBR 14724:2024, seções 5.1 a 5.9)
# ------------------------------------------------------------------
FONTE = "Arial"
TAM_TEXTO = 12          # 5.1 — fonte 12 para TODO o texto, inclusive a capa
TAM_MENOR = 10          # exceções: citações longas, notas, paginação, fontes/legendas
ESPACAMENTO = 1.5       # 5.2 — espaçamento 1,5 no corpo do texto
RECUO_PARAGRAFO = Cm(1.25)
RECUO_CITACAO = Cm(4)   # 5.2 / NBR 10520 — 4 cm da margem esquerda
MEIO_MANCHA = Cm(8)     # 5.2 — "do meio da mancha gráfica até a margem direita"
LARGURA_MANCHA = Cm(16) # 21 - 3 - 2 = 16 cm
MARGEM_ESQ, MARGEM_DIR, MARGEM_SUP, MARGEM_INF = Cm(3), Cm(2), Cm(3), Cm(2)
DISTANCIA_CABECALHO = Cm(2)   # 5.3 — número a 2 cm da borda superior

COR_PENDENCIA = RGBColor(0xC0, 0x00, 0x00)
COR_PRETO = RGBColor(0x00, 0x00, 0x00)

# Elementos sem indicativo numérico que ENTRAM no sumário (NBR 6027)
ELEMENTOS_NO_SUMARIO = (
    "REFERÊNCIAS", "REFERENCIAS", "APÊNDICE", "APENDICE",
    "ANEXO", "GLOSSÁRIO", "GLOSSARIO", "ÍNDICE", "INDICE",
)

RAIZ = Path(__file__).resolve().parents[2]
ORIGEM_PADRAO = RAIZ / "docs" / "tcc" / "TCC-Vida+.md"
DESTINO_PADRAO = RAIZ / "docs" / "tcc" / "TCC-Vida+.docx"


# ==================================================================
# Utilidades de baixo nível (OOXML)
# ==================================================================
def _ordem_de(parent: str) -> list:
    """Ordem canônica dos filhos de alguns elementos do OOXML."""
    return {
        "rPr": [
            "w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps",
            "w:smallCaps", "w:strike", "w:dstrike", "w:outline", "w:shadow",
            "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish",
            "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern",
            "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect",
            "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs",
            "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath",
        ],
        "pPr": [
            "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore",
            "w:framePr", "w:widowControl", "w:numPr", "w:suppressLineNumbers",
            "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
            "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
            "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
            "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
            "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
            "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
            "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
        ],
        "tblPr": [
            "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual",
            "w:tblStyleRowBandSize", "w:tblStyleColBandSize", "w:tblW",
            "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
            "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
            "w:tblDescription", "w:tblPrChange",
        ],
        "sectPr": [
            "w:footnotePr", "w:endnotePr", "w:type", "w:pgSz", "w:pgMar",
            "w:paperSrc", "w:pgBorders", "w:lnNumType", "w:pgNumType",
            "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
            "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
            "w:printerSettings", "w:sectPrChange",
        ],
    }[parent]


def inserir_em_ordem(parent, elemento, nome_parent: str):
    """Insere `elemento` respeitando a ordem canônica dos filhos de parent."""
    ordem = _ordem_de(nome_parent)
    nome_elemento = "w:" + elemento.tag.split("}")[-1]
    for filho in parent:
        nome_filho = "w:" + filho.tag.split("}")[-1]
        if nome_filho in ordem and ordem.index(nome_filho) > ordem.index(nome_elemento):
            filho.addprevious(elemento)
            return
    parent.append(elemento)


def fonte_run(run, tamanho=None, negrito=None, italico=None, cor=None, fonte=FONTE):
    run.font.name = fonte
    if tamanho is not None:
        run.font.size = Pt(tamanho)
    if negrito is not None:
        run.font.bold = negrito
    if italico is not None:
        run.font.italic = italico
    if cor is not None:
        run.font.color.rgb = cor
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(1 if rPr.find(qn("w:rStyle")) is not None else 0, rFonts)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(atributo), fonte)
    return run


def fonte_estilo(style, tamanho=None, negrito=None, italico=None, cor=None, fonte=FONTE):
    style.font.name = fonte
    if tamanho is not None:
        style.font.size = Pt(tamanho)
    if negrito is not None:
        style.font.bold = negrito
    if italico is not None:
        style.font.italic = italico
    if cor is not None:
        style.font.color.rgb = cor
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(1 if rPr.find(qn("w:rStyle")) is not None else 0, rFonts)
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(atributo), fonte)


def campal(campo: str) -> tuple:
    """Retorna os elementos <w:fldChar begin>, <w:instrText> e <w:fldChar end>."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = campo
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    return begin, instr, end


# ==================================================================
# Configuração do documento (estilos, página, seções)
# ==================================================================
def configurar_estilos(doc: Document):
    normal = doc.styles["Normal"]
    fonte_estilo(normal, tamanho=TAM_TEXTO, cor=COR_PRETO)
    pf = normal.paragraph_format
    pf.line_spacing = ESPACAMENTO
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = RECUO_PARAGRAFO
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    # 5.2.2 — títulos separados do texto por um espaço de 1,5 entre linhas
    titulos = {
        "Heading 1": dict(negrito=True, italico=False, centralizado=False),
        "Heading 2": dict(negrito=True, italico=False, centralizado=False),
        "Heading 3": dict(negrito=False, italico=True, centralizado=False),
        "Heading 4": dict(negrito=True, italico=True, centralizado=False),
    }
    for nome, cfg in titulos.items():
        try:
            style = doc.styles[nome]
        except KeyError:  # template sem o estilo: cria
            style = doc.styles.add_style(nome, WD_STYLE_TYPE.PARAGRAPH)
        fonte_estilo(style, tamanho=TAM_TEXTO, negrito=cfg["negrito"],
                     italico=cfg["italico"], cor=COR_PRETO)
        pf = style.paragraph_format
        pf.line_spacing = ESPACAMENTO
        pf.space_before = Pt(0)
        pf.space_after = Pt(18)          # ≈ uma linha de 1,5
        pf.first_line_indent = Cm(0)
        pf.keep_with_next = True
        pf.alignment = (WD_ALIGN_PARAGRAPH.CENTER if cfg["centralizado"]
                        else WD_ALIGN_PARAGRAPH.LEFT)

    # Títulos SEM indicativo numérico que não entram no sumário
    if "ElementoPreTextual" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("ElementoPreTextual", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["ElementoPreTextual"]
    fonte_estilo(style, tamanho=TAM_TEXTO, negrito=True, italico=False, cor=COR_PRETO)
    pf = style.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER   # 5.2.3 — centralizados
    pf.line_spacing = ESPACAMENTO
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(18)
    pf.keep_with_next = True

    # Texto de tabela (IBGE / NBR 14724 5.9) — fonte menor uniforme
    if "TextoTabela" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("TextoTabela", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["TextoTabela"]
    fonte_estilo(style, tamanho=TAM_MENOR, cor=COR_PRETO)
    pf = style.paragraph_format
    pf.line_spacing = 1.0
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


def configurar_secao(secao, cabecalho_paginado: bool = False, inicio: int | None = None):
    secao.page_width = Cm(21)      # 5.1 — A4
    secao.page_height = Cm(29.7)
    secao.left_margin = MARGEM_ESQ
    secao.top_margin = MARGEM_SUP
    secao.right_margin = MARGEM_DIR
    secao.bottom_margin = MARGEM_INF
    secao.header_distance = DISTANCIA_CABECALHO
    secao.footer_distance = Cm(2)

    if inicio is not None:
        pgNumType = OxmlElement("w:pgNumType")
        pgNumType.set(qn("w:start"), str(inicio))
        inserir_em_ordem(secao._sectPr, pgNumType, "sectPr")

    if not cabecalho_paginado:
        return
    secao.header.is_linked_to_previous = False
    par = secao.header.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pf = par.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.line_spacing = 1.0
    pf.space_after = Pt(0)
    run = par.add_run()
    begin, instr, end = campal(" PAGE ")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    fonte_run(run, tamanho=TAM_MENOR, cor=COR_PRETO)


# ==================================================================
# Leitura do Markdown
# ==================================================================
INLINE = re.compile(r"(\[\[.+?\]\]|\*\*.+?\*\*|\*[^*\n]+?\*)")
ETIQUETA = re.compile(r"^\[(c|cb|r|rb|n|e|l|s|t|f|a)(\d+)?\]\s?(.*)$", re.S)


def _limpar_inline(texto: str) -> str:
    texto = re.sub(r"\[\[(.+?)\]\]", r"\1", texto)
    texto = texto.replace("**", "").replace("*", "")
    return texto


def ler_md(caminho: Path) -> list:
    linhas = caminho.read_text(encoding="utf-8").splitlines()
    blocos, i = [], 0
    while i < len(linhas):
        linha = linhas[i]
        limpa = linha.strip()

        if not limpa:
            i += 1
            continue

        # comentários HTML (não vão para o .docx)
        if limpa.startswith("<!--"):
            conteudo = limpa
            # o comentário só termina na linha que TERMINA com "-->"
            while not conteudo.strip().endswith("-->") and i + 1 < len(linhas):
                i += 1
                conteudo += " " + linhas[i].strip()
            interno = (conteudo.replace("<!--", " ").replace("-->", " ")
                       .strip().upper().replace("Á", "A"))
            if interno == "SUMARIO":   # marcador <!-- SUMARIO --> = sumário automático
                blocos.append(("sumario", None))
            i += 1
            continue

        # títulos
        m = re.match(r"^(#{1,4})\s+(.*)$", limpa)
        if m:
            blocos.append(("h", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue

        # tabelas
        if limpa.startswith("|"):
            grade = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                celulas = [c.strip() for c in linhas[i].strip().strip("|").split("|")]
                grade.append(celulas)
                i += 1
            grade = [linha for linha in grade
                     if not all(re.fullmatch(r":?-{2,}:?", c or "-") for c in linha)]
            blocos.append(("tabela", grade))
            continue

        # citações com mais de três linhas
        if limpa.startswith(">"):
            texto = []
            while i < len(linhas) and linhas[i].strip().startswith(">"):
                texto.append(linhas[i].strip()[1:].strip())
                i += 1
            blocos.append(("citacao", " ".join(t for t in texto if t)))
            continue

        # ilustrações
        m = re.match(r"^!\[(.*?)\]\((.*?)\)\s*$", limpa)
        if m:
            blocos.append(("imagem", m.group(1), m.group(2)))
            i += 1
            continue

        # parágrafos com etiqueta de posicionamento
        m = ETIQUETA.match(limpa)
        if m:
            blocos.append(("tag", m.group(1), m.group(2), m.group(3)))
            i += 1
            continue

        blocos.append(("p", limpa))
        i += 1
    return blocos


# ==================================================================
# Escrita dos blocos no documento
# ==================================================================
def escrever_inline(paragrafo, texto: str, tamanho: int = TAM_TEXTO):
    for pedaco in INLINE.split(texto):
        if not pedaco:
            continue
        if pedaco.startswith("[[") and pedaco.endswith("]]"):
            fonte_run(paragrafo.add_run(pedaco[2:-2]), tamanho=tamanho,
                      italico=True, cor=COR_PENDENCIA)
        elif pedaco.startswith("**") and pedaco.endswith("**"):
            fonte_run(paragrafo.add_run(pedaco[2:-2]), tamanho=tamanho, negrito=True,
                      cor=COR_PRETO)
        elif pedaco.startswith("*") and pedaco.endswith("*"):
            fonte_run(paragrafo.add_run(pedaco[1:-1]), tamanho=tamanho, italico=True,
                      cor=COR_PRETO)
        else:
            fonte_run(paragrafo.add_run(pedaco), tamanho=tamanho, cor=COR_PRETO)
    return paragrafo


def paragrafo_vazio(doc, linhas: int = 1, tamanho: int = TAM_TEXTO):
    for _ in range(linhas):
        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.line_spacing = 1.0
        pf.space_after = Pt(0)
        fonte_run(par.add_run(""), tamanho=tamanho)


def escrever_sumario(doc, entradas):
    """Campo TOC do Word com resultado em cache (atualize no Word/Docs)."""
    if not entradas:
        return
    primeiro = doc.add_paragraph()
    pf = primeiro.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.line_spacing = ESPACAMENTO
    pf.space_after = Pt(0)
    run = primeiro.add_run()
    begin, instr, _ = campal(' TOC \\o "1-3" \\h \\z \\u ')
    run._r.append(begin)
    run._r.append(instr)
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    run._r.append(separador)

    for nivel, texto in entradas:
        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.line_spacing = ESPACAMENTO
        pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # sumário em caixa alta apenas na seção primária (NBR 6027)
        par.paragraph_format.tab_stops.add_tab_stop(
            LARGURA_MANCHA, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if nivel == 1:
            escrever_inline(par, "**" + texto.upper() + "**")
        elif nivel == 2:
            par.paragraph_format.left_indent = Cm(0.5)
            escrever_inline(par, "**" + texto + "**")
        else:
            par.paragraph_format.left_indent = Cm(1.0)
            escrever_inline(par, "*" + texto + "*")
        par.add_run("\t")

    fim = doc.add_paragraph()
    pf = fim.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.line_spacing = 1.0
    pf.space_after = Pt(0)
    run = fim.add_run()
    run._r.append(OxmlElement("w:fldChar"))
    run._r[-1].set(qn("w:fldCharType"), "end")


def escrever_tabela(doc, grade):
    if not grade:
        return
    colunas = max(len(linha) for linha in grade)
    tabela = doc.add_table(rows=0, cols=colunas)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.autofit = True

    for indice, linha in enumerate(grade):
        celulas = tabela.add_row().cells
        for posicao in range(colunas):
            texto = _limpar_inline(linha[posicao]) if posicao < len(linha) else ""
            par = celulas[posicao].paragraphs[0]
            par.style = doc.styles["TextoTabela"]
            par.alignment = (WD_ALIGN_PARAGRAPH.CENTER if posicao or indice == 0
                             else WD_ALIGN_PARAGRAPH.LEFT)
            if indice == 0:
                fonte_run(par.add_run(texto), tamanho=TAM_MENOR, negrito=True,
                          cor=COR_PRETO)
            else:
                fonte_run(par.add_run(texto), tamanho=TAM_MENOR, cor=COR_PRETO)
        if indice == 0:  # filete sob o cabeçalho (NBR 14724 5.9 / IBGE)
            for celula in celulas:
                tcPr = celula._tc.get_or_add_tcPr()
                bordas = OxmlElement("w:tcBorders")
                fundo = OxmlElement("w:bottom")
                fundo.set(qn("w:val"), "single")
                fundo.set(qn("w:sz"), "6")
                fundo.set(qn("w:color"), "000000")
                bordas.append(fundo)
                tcPr.append(bordas)

    # tabela "aberta": apenas filetes superior e inferior
    tblPr = tabela._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for aresta, tamanho, estilo in (("top", "12", "single"), ("bottom", "12", "single"),
                                    ("left", "0", "none"), ("right", "0", "none"),
                                    ("insideH", "0", "none"), ("insideV", "0", "none")):
        elemento = OxmlElement("w:" + aresta)
        elemento.set(qn("w:val"), estilo)
        elemento.set(qn("w:sz"), tamanho)
        elemento.set(qn("w:space"), "0")
        elemento.set(qn("w:color"), "000000")
        bordas.append(elemento)
    inserir_em_ordem(tblPr, bordas, "tblPr")

    # respiro depois da tabela
    par = doc.add_paragraph()
    par.paragraph_format.first_line_indent = Cm(0)
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.line_spacing = 1.0
    fonte_run(par.add_run(""), tamanho=TAM_MENOR)


def construir(blocos, destino: Path):
    doc = Document()
    configurar_estilos(doc)
    configurar_secao(doc.sections[0])

    entradas_sumario = []
    for bloco in blocos:
        if bloco[0] == "h":
            _, nivel, texto = bloco
            numerado = texto[:1].isdigit()
            if nivel == 1 and (numerado or texto.upper().startswith(ELEMENTOS_NO_SUMARIO)):
                entradas_sumario.append((nivel, _limpar_inline(texto)))
            elif nivel > 1 and entradas_sumario:
                entradas_sumario.append((nivel, _limpar_inline(texto)))

    em_referencias = False

    # primeira seção primária NUMERADA = início dos elementos textuais
    indice_inicio_textual = next(
        (i for i, b in enumerate(blocos)
         if b[0] == "h" and b[1] == 1 and b[2][:1].isdigit()),
        None,
    )
    # 5.3 — folhas pré-textuais são contadas (a capa não) mas não numeradas
    paginas_pre_textuais = (
        sum(1 for i, b in enumerate(blocos)
            if i < indice_inicio_textual and b[0] == "h" and b[1] == 1)
        if indice_inicio_textual is not None else 0
    )

    for indice, bloco in enumerate(blocos):
        tipo = bloco[0]

        if tipo == "h":
            _, nivel, texto = bloco
            numerado = texto[:1].isdigit()
            entra_sumario = (nivel == 1 and
                             (numerado or texto.upper().startswith(ELEMENTOS_NO_SUMARIO)))
            em_referencias = texto.upper().startswith(("REFERÊNCIAS", "REFERENCIAS"))
            abre_secao_textual = (indice == indice_inicio_textual)

            if abre_secao_textual:
                nova = doc.add_section(WD_SECTION.NEW_PAGE)
                configurar_secao(nova, cabecalho_paginado=True,
                                 inicio=max(paginas_pre_textuais, 1))

            if nivel == 1:
                par = doc.add_paragraph(style="Heading 1" if entra_sumario
                                        else "ElementoPreTextual")
                par.alignment = (WD_ALIGN_PARAGRAPH.LEFT
                                 if (entra_sumario and numerado)
                                 else WD_ALIGN_PARAGRAPH.CENTER)
                # 5.2.2 — cada seção primária começa em nova página
                if indice > 0 and not abre_secao_textual:
                    par.paragraph_format.page_break_before = True
            else:
                par = doc.add_paragraph(style=f"Heading {min(nivel, 4)}")
            escrever_inline(par, texto)
            continue

        if tipo == "sumario":
            escrever_sumario(doc, entradas_sumario)
            continue

        if tipo == "p":
            par = doc.add_paragraph()
            if em_referencias:  # 5.2 — referências em espaço simples
                pf = par.paragraph_format
                pf.first_line_indent = Cm(0)
                pf.line_spacing = 1.0
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf.space_after = Pt(14)   # "espaço simples em branco" entre referências
            escrever_inline(par, bloco[1])
            continue

        if tipo == "citacao":
            par = doc.add_paragraph()
            pf = par.paragraph_format
            pf.left_indent = RECUO_CITACAO
            pf.first_line_indent = Cm(0)
            pf.line_spacing = 1.0
            pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            par.add_run("")
            escrever_inline(par, bloco[1], tamanho=TAM_MENOR)
            continue

        if tipo == "tag":
            _, etiqueta, tamanho, texto = bloco
            if etiqueta == "e":
                paragrafo_vazio(doc)
                continue
            par = doc.add_paragraph()
            pf = par.paragraph_format
            pf.first_line_indent = Cm(0)
            tamanho_fonte = int(tamanho) if tamanho else TAM_TEXTO
            if etiqueta == "n":          # natureza / dedicatória / epígrafe (5.2)
                pf.left_indent = MEIO_MANCHA
                pf.line_spacing = 1.0
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            elif etiqueta == "s":        # resumo, abstract, palavras-chave
                pf.line_spacing = ESPACAMENTO
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            elif etiqueta == "l":        # lista de ilustrações, tabelas, siglas
                pf.line_spacing = ESPACAMENTO
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf.tab_stops.add_tab_stop(LARGURA_MANCHA, WD_TAB_ALIGNMENT.RIGHT,
                                          WD_TAB_LEADER.DOTS)
            elif etiqueta == "t":        # título de ilustração ou tabela (5.2)
                pf.line_spacing = 1.0
                pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf.space_after = Pt(4)
            elif etiqueta == "a":        # alínea: a) b) c) — NBR 6024:2012
                pf.line_spacing = ESPACAMENTO
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                pf.left_indent = Cm(1.25)
            elif etiqueta == "f":        # fonte/legenda de ilustração ou tabela
                pf.line_spacing = 1.0
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf.space_after = Pt(12)
                tamanho_fonte = TAM_MENOR
            else:
                pf.line_spacing = ESPACAMENTO
                pf.alignment = {
                    "c": WD_ALIGN_PARAGRAPH.CENTER,
                    "cb": WD_ALIGN_PARAGRAPH.CENTER,
                    "r": WD_ALIGN_PARAGRAPH.RIGHT,
                    "rb": WD_ALIGN_PARAGRAPH.RIGHT,
                }[etiqueta]
            if etiqueta == "l" and "|" in texto:
                item, pagina = texto.split("|", 1)
                escrever_inline(par, item.strip(), tamanho=tamanho_fonte)
                par.add_run("\t")
                escrever_inline(par, pagina.strip(), tamanho=tamanho_fonte)
                continue
            if etiqueta in ("cb", "rb") or (etiqueta == "t" and "**" not in texto):
                escrever_inline(par, "**" + texto + "**", tamanho=tamanho_fonte)
            else:
                escrever_inline(par, texto, tamanho=tamanho_fonte)
            continue

        if tipo == "imagem":
            _, titulo, caminho = bloco
            arquivo = (RAIZ / caminho).resolve()
            if not arquivo.exists():
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.first_line_indent = Cm(0)
                escrever_inline(par, f"[[ILUSTRAÇÃO A INSERIR: {titulo or caminho}]]")
                continue
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.first_line_indent = Cm(0)
            par.paragraph_format.line_spacing = 1.0
            par.add_run().add_picture(str(arquivo), width=Cm(15))
            continue

        if tipo == "tabela":
            escrever_tabela(doc, bloco[1])
            continue

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
    return doc


def main(argv: list) -> int:
    origem = Path(argv[1]).resolve() if len(argv) > 1 else ORIGEM_PADRAO
    destino = Path(argv[2]).resolve() if len(argv) > 2 else DESTINO_PADRAO
    if not origem.exists():
        print(f"[erro] arquivo não encontrado: {origem}")
        return 1
    blocos = ler_md(origem)
    doc = construir(blocos, destino)

    pendentes = sum(1 for p in doc.paragraphs for r in p.runs
                    if r.font.color and r.font.color.rgb == COR_PENDENCIA)
    faltando = sum(1 for p in doc.paragraphs for r in p.runs
                   if r.text.strip().startswith("ILUSTRAÇÃO A INSERIR"))
    indice = next((i for i, b in enumerate(blocos)
                   if b[0] == "h" and b[1] == 1 and b[2][:1].isdigit()), None)
    pre_textuais = sum(1 for i, b in enumerate(blocos)
                       if indice is not None and i < indice and b[0] == "h" and b[1] == 1)
    primarias = sum(1 for b in blocos if b[0] == "h" and b[1] == 1)
    print(f"[ok] documento gerado: {destino}")
    print(f"     seções primárias no total: {primarias}")
    print(f"     páginas pré-textuais contadas: {max(pre_textuais - 1, 0)} "
          f"(a capa não é contada) → numeração inicia na folha {pre_textuais}")
    print(f"     pendências em vermelho: {pendentes}")
    if faltando:
        print(f"     ilustrações a inserir: {faltando}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
